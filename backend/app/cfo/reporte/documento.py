"""FABS · ensambla el .docx del reporte a inversionistas (python-docx).
PURO: recibe DatosReporte + PNG y devuelve los bytes del documento; no lee
servicios ni dominio (S1)."""

from io import BytesIO

from docx import Document
from docx.shared import Inches

from app.cfo.reporte.modelos import DatosReporte


def armar_docx(datos: DatosReporte, png_rumbo: bytes | None) -> bytes:
    doc = Document()
    titulo = f"RODDOS S.A.S. — Reporte a inversionistas — {datos.periodo}"
    doc.add_heading(titulo, level=0)
    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(datos.resumen_ejecutivo)
    if png_rumbo:
        doc.add_picture(BytesIO(png_rumbo), width=Inches(6.0))
    for sec in datos.secciones:
        doc.add_heading(sec.titulo, level=1)
        if sec.cifras:
            tabla = doc.add_table(rows=0, cols=2)
            tabla.style = "Light Grid Accent 1"
            for c in sec.cifras:
                fila = tabla.add_row().cells
                fila[0].text = c.etiqueta
                fila[1].text = c.valor
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
