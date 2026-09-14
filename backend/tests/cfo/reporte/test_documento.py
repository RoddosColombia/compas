import base64
from io import BytesIO

from app.cfo.reporte.documento import armar_docx
from app.cfo.reporte.modelos import Cifra, DatosReporte, RumboSerie, Seccion
from docx import Document


def _datos():
    resumen = (
        "La caja disponible es $5.000.000 y el piso proyectado cruza el "
        "umbral en abril-2027."
    )
    return DatosReporte(
        periodo="2026-09",
        resumen_ejecutivo=resumen,
        secciones=[
            Seccion(titulo="Caja y rumbo al umbral", cifras=[
                Cifra(etiqueta="Caja disponible hoy", valor="$5.000.000"),
                Cifra(etiqueta="Piso proyectado", valor="$1.200.000"),
            ]),
            Seccion(titulo="Ventas / colocación", cifras=[
                Cifra(etiqueta="Mix Raider", valor="45,0%"),
            ]),
        ],
        rumbo=RumboSerie(meses=["2026-09", "2026-10"], caja=["5000000", "4200000"],
                         caja_minima="1000000", caja_atencion="3000000"),
    )


def test_armar_docx_valido_con_secciones_y_grafica():
    # sin imagen para no exigir un PNG válido en este caso
    data = armar_docx(_datos(), png_rumbo=None)
    assert isinstance(data, bytes) and len(data) > 0
    doc = Document(BytesIO(data))  # round-trip: abre
    textos = [p.text for p in doc.paragraphs]
    assert any("Reporte a inversionistas" in t for t in textos)
    assert any("Resumen ejecutivo" in t for t in textos)
    assert any("Caja y rumbo al umbral" in t for t in textos)
    # las cifras van en tablas
    celdas = [c.text for tbl in doc.tables for row in tbl.rows for c in row.cells]
    assert "Caja disponible hoy" in celdas and "$5.000.000" in celdas


def test_armar_docx_embebe_la_grafica_si_hay_png(tmp_path):
    # PNG real mínimo 1x1 (para que add_picture no falle). NOTA: el literal
    # base64 sugerido originalmente en el brief está truncado (IDAT desborda
    # el buffer, sin IEND válido) y python-docx lo rechaza con
    # UnexpectedEndOfFileError — verificado leyendo los chunks a mano. Este
    # literal se generó con Pillow
    # (Image.new('RGB', (1, 1)).save(..., format='PNG')) y decodifica a un
    # PNG válido de 69 bytes.
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4z8AA"
        "AAMBAQDJ/pLvAAAAAElFTkSuQmCC"
    )
    data = armar_docx(_datos(), png_rumbo=png)
    doc = Document(BytesIO(data))
    assert len(doc.inline_shapes) >= 1  # la imagen quedó embebida
